import io
import csv
import datetime
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from backend.app.utils.sanitizer import format_currency_vnd

class ReportService:
    """
    Generates downloadable reports in Excel (.xlsx), CSV (.csv), and PDF (.pdf) formats.
    """

    @staticmethod
    def generate_excel_report(
        user_name: str,
        transactions: List[Dict[str, Any]],
        summary_data: Dict[str, Any]
    ) -> io.BytesIO:
        """Generates a professional Excel (.xlsx) financial workbook."""
        wb = Workbook()
        
        # 1. Summary Sheet
        ws_sum = wb.active
        ws_sum.title = "Tổng Quan Tài Chính"
        ws_sum.views.sheetView[0].showGridLines = True

        # Header Title
        ws_sum.merge_cells("A1:E1")
        title_cell = ws_sum["A1"]
        title_cell.value = f"BÁO CÁO TÀI CHÍNH CÁ NHÂN - FINTRACK AI"
        title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_sum.row_dimensions[1].height = 40

        ws_sum["A2"] = f"Người dùng: {user_name}"
        ws_sum["A3"] = f"Ngày xuất báo cáo: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}"
        for r in range(2, 4):
            ws_sum[f"A{r}"].font = Font(italic=True, size=11)

        # Summary KPIs Table
        ws_sum["A5"] = "CHỈ SỐ TÀI CHÍNH"
        ws_sum["B5"] = "GIÁ TRỊ"
        ws_sum["A5"].font = Font(bold=True, color="FFFFFF")
        ws_sum["B5"].font = Font(bold=True, color="FFFFFF")
        ws_sum["A5"].fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        ws_sum["B5"].fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")

        kpi_rows = [
            ("Tổng Thu Nhập", summary_data.get("total_income", 0)),
            ("Tổng Chi Tiêu", summary_data.get("total_expense", 0)),
            ("Dòng Tiền Ròng (Net Flow)", summary_data.get("net_savings", 0)),
            ("Tỷ Lệ Tiết Kiệm", f"{summary_data.get('savings_rate', 0)}%"),
            ("Tổng Số Giao Dịch", len(transactions))
        ]

        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        for idx, (label, val) in enumerate(kpi_rows, start=6):
            ws_sum[f"A{idx}"] = label
            ws_sum[f"B{idx}"] = val
            ws_sum[f"A{idx}"].border = thin_border
            ws_sum[f"B{idx}"].border = thin_border
            if isinstance(val, (int, float)):
                ws_sum[f"B{idx}"].number_format = '#,##0 "₫"'

        ws_sum.column_dimensions["A"].width = 30
        ws_sum.column_dimensions["B"].width = 25

        # 2. Transaction Details Sheet
        ws_tx = wb.create_sheet(title="Chi Tiết Giao Dịch")
        ws_tx.views.sheetView[0].showGridLines = True

        headers = ["STT", "Thời Gian", "Loại Giao Dịch", "Số Tiền", "Danh Mục", "Ví Thanh Toán", "Ghi Chú", "Nguồn Tạo"]
        ws_tx.append(headers)

        for col_num, h_name in enumerate(headers, 1):
            c = ws_tx.cell(row=1, column=col_num)
            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
        ws_tx.row_dimensions[1].height = 28

        for idx, t in enumerate(transactions, start=1):
            tx_type_vn = "Chi tiêu" if t.get("type") == "EXPENSE" else "Thu nhập" if t.get("type") == "INCOME" else "Chuyển tiền"
            row_data = [
                idx,
                t.get("date", ""),
                tx_type_vn,
                t.get("amount", 0),
                t.get("category_name", ""),
                t.get("wallet_name", ""),
                t.get("note", ""),
                "AI Nhập" if t.get("created_by_ai") == "AI_PARSED" else "Thủ công"
            ]
            ws_tx.append(row_data)
            row_idx = idx + 1
            ws_tx.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(row_data) + 1):
                cell = ws_tx.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if col_idx == 4:  # Amount column
                    cell.number_format = '#,##0 "₫"'
                    if t.get("type") == "EXPENSE":
                        cell.font = Font(color="DC2626")
                    elif t.get("type") == "INCOME":
                        cell.font = Font(color="16A34A")

        # Auto fit column widths
        for col in ws_tx.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws_tx.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generate_csv_report(transactions: List[Dict[str, Any]]) -> str:
        """Generates CSV with UTF-8 BOM for Windows Excel compatibility."""
        output = io.StringIO()
        # UTF-8 BOM
        output.write('\ufeff')
        writer = csv.writer(output, delimiter=',', quoting=csv.QUOTE_MINIMAL)
        writer.writerow(["ID", "Ngay", "Loai", "So_Tien_VND", "Danh_Muc", "Vi_Thanh_Toan", "Ghi_Chu", "Nguon"])

        for t in transactions:
            tx_type_vn = "EXPENSE" if t.get("type") == "EXPENSE" else "INCOME" if t.get("type") == "INCOME" else "TRANSFER"
            writer.writerow([
                t.get("id", ""),
                t.get("date", ""),
                tx_type_vn,
                t.get("amount", 0),
                t.get("category_name", ""),
                t.get("wallet_name", ""),
                t.get("note", "").replace("\n", " "),
                t.get("created_by_ai", "MANUAL")
            ])

        return output.getvalue()

    @staticmethod
    def generate_pdf_report(
        user_name: str,
        transactions: List[Dict[str, Any]],
        summary_data: Dict[str, Any]
    ) -> io.BytesIO:
        """Generates clean, styled PDF report."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
        elements = []
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#10B981"),
            alignment=1,  # Center
            spaceAfter=12
        )
        subtitle_style = ParagraphStyle(
            name="SubTitleStyle",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#6B7280"),
            alignment=1,
            spaceAfter=20
        )

        elements.append(Paragraph("FINTRACK AI - BAO CAO TAI CHINH CA NHAN", title_style))
        elements.append(Paragraph(f"Nguoi dung: {user_name} | Ngay xuat: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))

        # KPI Summary Table
        kpi_data = [
            ["Chi so", "Gia tri"],
            ["Tong Thu Nhap", format_currency_vnd(summary_data.get("total_income", 0))],
            ["Tong Chi Tieu", format_currency_vnd(summary_data.get("total_expense", 0))],
            ["Tiet Kiem Rong", format_currency_vnd(summary_data.get("net_savings", 0))],
            ["Ty Le Tiet Kiem", f"{summary_data.get('savings_rate', 0)}%"],
            ["Tong So Giao Dich", str(len(transactions))]
        ]

        t_summary = Table(kpi_data, colWidths=[200, 250])
        t_summary.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#10B981')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')])
        ]))
        elements.append(t_summary)
        elements.append(Spacer(1, 20))

        # Recent Transactions Header
        section_style = ParagraphStyle(
            name="SectionStyle",
            parent=styles["Heading2"],
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#1F2937"),
            spaceAfter=10
        )
        elements.append(Paragraph("Danh Sach Giao Dich (Toi da 25 giao dich gan nhat)", section_style))

        tx_rows = [["Ngay", "Loai", "So Tien", "Danh Muc", "Vi", "Ghi Chu"]]
        for t in transactions[:25]:
            tx_type = "Chi" if t.get("type") == "EXPENSE" else "Thu" if t.get("type") == "INCOME" else "Chuyen"
            amt_str = format_currency_vnd(t.get("amount", 0))
            tx_rows.append([
                str(t.get("date", ""))[:10],
                tx_type,
                amt_str,
                str(t.get("category_name", ""))[:15],
                str(t.get("wallet_name", ""))[:12],
                str(t.get("note", ""))[:20]
            ])

        t_tx = Table(tx_rows, colWidths=[65, 45, 90, 100, 85, 140])
        t_tx.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#059669')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')])
        ]))
        elements.append(t_tx)

        doc.build(elements)
        buffer.seek(0)
        return buffer

report_service = ReportService()
