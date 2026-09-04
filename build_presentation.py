"""
Script tạo Slide Báo Cáo Đề Tài "FinTrack AI - Ứng Dụng Quản Lý Chi Tiêu Cá Nhân Tích Hợp Trí Tuệ Nhân Tạo"
Nhóm 03: Đặng Quyết Thắng (Trưởng nhóm), Nguyễn Văn Tiến, Quách Minh Hiếu
Giảng viên hướng dẫn: ThS. Hà Thị Thanh (ICTU)
Phong cách: Dark Cyber Neon Theme (Glassmorphism, 16:9 Widescreen, Tiếng Việt tự nhiên, dễ thuyết trình)
"""

import sys
import os

# Cấu hình encoding stdout để tránh lỗi charmap trên Windows console
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# BẢNG MÀU CHỦ ĐẠO (DARK CYBER NEON THEME)
# ==============================================================================
COLOR_BG_DARK = RGBColor(11, 15, 25)       # #0B0F19 Nền tối chủ đạo
COLOR_BG_CARD = RGBColor(17, 24, 39)       # #111827 Nền thẻ Card Glassmorphism
COLOR_BG_CARD_LIGHT = RGBColor(24, 33, 53) # #182135 Nền thẻ phụ
COLOR_BG_BADGE = RGBColor(30, 41, 59)      # #1E293B Nền Badge tag

# Neon Accent Colors
COLOR_NEON_CYAN = RGBColor(0, 229, 255)    # #00E5FF Xanh ngọc Neon
COLOR_NEON_EMERALD = RGBColor(0, 255, 170) # #00FFAA Xanh lá Neon
COLOR_NEON_AMBER = RGBColor(245, 158, 11)  # #F59E0B Vàng cam Gold
COLOR_NEON_PINK = RGBColor(255, 0, 122)    # #FF007A Hồng Neon
COLOR_NEON_PURPLE = RGBColor(139, 92, 246) # #8B5CF6 Tím Neon
COLOR_NEON_BLUE = RGBColor(59, 130, 246)   # #3B82F6 Xanh dương

# Text Colors
COLOR_TEXT_WHITE = RGBColor(255, 255, 255) # Trắng sáng
COLOR_TEXT_LIGHT = RGBColor(226, 232, 240) # Trắng xám nhạt (dễ đọc)
COLOR_TEXT_MUTED = RGBColor(148, 163, 184) # Xám vừa
COLOR_TEXT_DIM = RGBColor(100, 116, 139)   # Xám mờ

FONT_MAIN = "Segoe UI"
FONT_HEADING = "Segoe UI"
TOTAL_SLIDES = 11

# ==============================================================================
# CÁC HÀM TIỆN ÍCH XÂY DỰNG GIAO DIỆN SLIDE
# ==============================================================================

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def set_slide_background(slide):
    """Vẽ nền tối toàn màn hình cho Slide"""
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG_DARK
    bg.line.fill.background()
    return bg

def add_header(slide, tag_text, title_text, subtitle_text, accent_color=COLOR_NEON_CYAN):
    """Tạo Header chuẩn hóa với Tag, Tiêu đề lớn và Phụ đề"""
    # 1. Tag pill
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.0), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.name = FONT_MAIN
    p_tag.font.size = Pt(10.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = accent_color

    # 2. Tiêu đề chính
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.55))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(21)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TEXT_WHITE

    # 3. Phụ đề tóm lược
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.25), Inches(11.7), Inches(0.35))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(11.5)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    # 4. Đường kẻ neon ngăn cách header
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.733), Inches(0.02)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = accent_color
    line.line.fill.background()

def add_footer(slide, slide_num, total_slides=TOTAL_SLIDES, accent_color=COLOR_NEON_CYAN):
    """Tạo Footer chuẩn hóa cho tất cả các slide nội dung"""
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.88), Inches(11.733), Inches(0.015)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(30, 41, 59)
    line.line.fill.background()

    # Text bên trái
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.96), Inches(7.0), Inches(0.35))
    tf = footer_box.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "✦ FinTrack AI • Ứng Dụng Quản Lý Chi Tiêu Cá Nhân Tích Hợp AI  |  Nhóm 03 (ICTU)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_DIM

    # Slide number bên phải
    num_box = slide.shapes.add_textbox(Inches(9.5), Inches(6.96), Inches(3.0), Inches(0.35))
    tf_num = num_box.text_frame
    tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
    p_num = tf_num.paragraphs[0]
    p_num.text = f"Slide {slide_num:02d} / {total_slides:02d}"
    p_num.alignment = PP_ALIGN.RIGHT
    p_num.font.name = FONT_MAIN
    p_num.font.size = Pt(9.5)
    p_num.font.bold = True
    p_num.font.color.rgb = accent_color

def add_card(slide, left, top, width, height, border_color=COLOR_NEON_CYAN, bg_color=COLOR_BG_CARD):
    """Tạo Card Glassmorphism viền Neon"""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    return card

# ==============================================================================
# HÀM TẠO 11 SLIDES THUYẾT TRÌNH DỄ HIỂU & DỄ NÓI
# ==============================================================================

def build_slide_1(prs):
    """SLIDE 1: TRANG BÌA"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    # Top Tag Ribbon
    top_badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.4), Inches(0.55), Inches(6.533), Inches(0.4)
    )
    top_badge.fill.solid()
    top_badge.fill.fore_color.rgb = COLOR_BG_BADGE
    top_badge.line.color.rgb = COLOR_NEON_CYAN
    top_badge.line.width = Pt(1)
    tf_badge = top_badge.text_frame
    p_b = tf_badge.paragraphs[0]
    p_b.text = "✦ BÁO CÁO ĐỀ TÀI HỌC PHẦN ỨNG DỤNG TRÍ TUỆ NHÂN TẠO ✦"
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.name = FONT_MAIN
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_NEON_CYAN

    # Main Giant Title
    main_title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.733), Inches(1.35))
    tf_title = main_title_box.text_frame
    tf_title.word_wrap = True
    
    p1 = tf_title.paragraphs[0]
    p1.text = "FINTRACK AI"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE

    p2 = tf_title.add_paragraph()
    p2.text = "ỨNG DỤNG QUẢN LÝ CHI TIÊU CÁ NHÂN TÍCH HỢP TRÍ TUỆ NHÂN TẠO"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(16.5)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_NEON_EMERALD

    p3 = tf_title.add_paragraph()
    p3.text = "Giải pháp ghi chép tài chính thông minh, bảo mật an toàn và nhắc nhở chi tiêu hiệu quả"
    p3.alignment = PP_ALIGN.CENTER
    p3.font.name = FONT_MAIN
    p3.font.size = Pt(11.5)
    p3.font.color.rgb = COLOR_TEXT_LIGHT

    # 3 Info Cards
    # Card 1: Học phần & Đơn vị
    add_card(slide, 0.8, 2.7, 3.65, 3.1, COLOR_NEON_CYAN)
    card1_tb = slide.shapes.add_textbox(Inches(0.95), Inches(2.85), Inches(3.35), Inches(2.8))
    tf1 = card1_tb.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "[ THÔNG TIN HỌC PHẦN ]"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_CYAN

    items1 = [
        ("Trường đào tạo:", "ĐH CNTT & Truyền Thông (ICTU)"),
        ("Khoa chuyên môn:", "Khoa Công Nghệ Thông Tin"),
        ("Học phần:", "Ứng Dụng Trí Tuệ Nhân Tạo"),
        ("Thời gian thực hiện:", "Năm 2026")
    ]
    for label, val in items1:
        p_l = tf1.add_paragraph()
        p_l.text = f"• {label} "
        p_l.font.size = Pt(10.5)
        p_l.font.bold = True
        p_l.font.color.rgb = COLOR_TEXT_LIGHT
        
        run_v = p_l.add_run()
        run_v.text = val
        run_v.font.bold = False
        run_v.font.color.rgb = COLOR_TEXT_MUTED

    # Card 2: Giảng viên hướng dẫn
    add_card(slide, 4.84, 2.7, 3.65, 3.1, COLOR_NEON_AMBER)
    card2_tb = slide.shapes.add_textbox(Inches(4.99), Inches(2.85), Inches(3.35), Inches(2.8))
    tf2 = card2_tb.text_frame
    tf2.word_wrap = True
    
    p = tf2.paragraphs[0]
    p.text = "[ GIẢNG VIÊN HƯỚNG DẪN ]"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_AMBER

    items2 = [
        ("Giảng viên hướng dẫn:", "ThS. Hà Thị Thanh"),
        ("Bộ môn:", "Trí Tuệ Nhân Tạo"),
        ("Vai trò:", "Định hướng & Hướng dẫn đề tài"),
        ("Đánh giá:", "Ứng dụng AI thiết thực cho sinh viên")
    ]
    for label, val in items2:
        p_l = tf2.add_paragraph()
        p_l.text = f"• {label} "
        p_l.font.size = Pt(10.5)
        p_l.font.bold = True
        p_l.font.color.rgb = COLOR_TEXT_LIGHT
        
        run_v = p_l.add_run()
        run_v.text = val
        run_v.font.bold = False
        run_v.font.color.rgb = COLOR_TEXT_MUTED

    # Card 3: Nhóm sinh viên thực hiện (Nhóm 03)
    add_card(slide, 8.88, 2.7, 3.65, 3.1, COLOR_NEON_EMERALD)
    card3_tb = slide.shapes.add_textbox(Inches(9.03), Inches(2.85), Inches(3.35), Inches(2.8))
    tf3 = card3_tb.text_frame
    tf3.word_wrap = True
    
    p = tf3.paragraphs[0]
    p.text = "[ NHÓM THỰC HIỆN - NHÓM 03 ]"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_EMERALD

    members = [
        ("1. Đặng Quyết Thắng", "Trưởng Nhóm / Phát Triển Chính & AI"),
        ("2. Nguyễn Văn Tiến", "Thành Viên / Thiết Kế Giao Diện Web"),
        ("3. Quách Minh Hiếu", "Thành Viên / Xử Lý Dữ Liệu & Backend")
    ]
    for name, role in members:
        p_m = tf3.add_paragraph()
        p_m.text = f"✦ {name}"
        p_m.font.size = Pt(10.5)
        p_m.font.bold = True
        p_m.font.color.rgb = COLOR_TEXT_WHITE
        
        p_r = tf3.add_paragraph()
        p_r.text = f"   ↳ {role}"
        p_r.font.size = Pt(9.5)
        p_r.font.color.rgb = COLOR_NEON_EMERALD if "Trưởng" in role else COLOR_TEXT_MUTED

    # Bottom Highlight Ribbon
    pills = [
        ("⚡ Nhập Chi Tiêu Bằng AI", COLOR_NEON_CYAN),
        ("🛡️ Bảo Mật Xóa Dữ Liệu Nhạy Cảm", COLOR_NEON_PINK),
        ("🩺 Cố Vấn Tiết Kiệm 50/30/20", COLOR_NEON_AMBER),
        ("🏦 Quét Mã VietQR Tự Động", COLOR_NEON_EMERALD),
        ("🎨 Giao Diện Đẹp Chuẩn 120 FPS", COLOR_NEON_PURPLE)
    ]
    total_w = 11.733
    pill_w = total_w / len(pills) - 0.1
    for i, (txt, col) in enumerate(pills):
        px = 0.8 + i * (pill_w + 0.1)
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(px), Inches(6.05), Inches(pill_w), Inches(0.48)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_BG_BADGE
        badge.line.color.rgb = col
        badge.line.width = Pt(1.2)
        tf_b = badge.text_frame
        p = tf_b.paragraphs[0]
        p.text = txt
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_MAIN
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_WHITE


def build_slide_2(prs):
    """SLIDE 2: VÌ SAO NHÓM CHỌN ĐỀ TÀI NÀY?"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(
        slide,
        "02 // LÝ DO CHỌN ĐỀ TÀI",
        "Vì Sao Nhóm Quyết Định Xây Dựng FinTrack AI?",
        "Giải quyết những khó khăn thực tế của giới trẻ và sinh viên trong việc quản lý tiền bạc hàng ngày"
    )

    # 4 Cards for 4 Problems / Goals
    cards = [
        ("💸 1. Tiêu Tiền Không Kiểm Soát",
         "Thói quen quét mã QR và chuyển khoản online quá tiện lợi khiến chúng ta dễ chi tiêu quá đà, cuối tháng hay bị rỗng ví mà không biết tiền đã đi đâu.",
         COLOR_NEON_PINK),
        ("📝 2. Lười Ghi Chép Sổ Sách",
         "Nhập liệu bằng tay quá mất thời gian và phiền phức. Hầu hết mọi người chỉ ghi chép được vài hôm rồi bỏ cuộc vì hay quên các khoản chi nhỏ.",
         COLOR_NEON_AMBER),
        ("🎯 3. Không Biết Cách Chia Tiền",
         "Thiếu kiến thức phân bổ thu nhập hợp lý (ăn uống bao nhiêu, tiết kiệm bao nhiêu), không có kế hoạch rõ ràng để dành tiền cho tương lai.",
         COLOR_NEON_PURPLE),
        ("💡 4. Mục Tiêu Của FinTrack AI",
         "Tạo ra trang web giúp ghi chép siêu nhanh chỉ bằng một câu nói tự nhiên, có AI nhắc nhở và hướng dẫn quản lý tài chính thông minh 24/7.",
         COLOR_NEON_EMERALD)
    ]

    card_w = 5.75
    card_h = 2.25
    coords = [
        (0.8, 1.85), (6.78, 1.85),
        (0.8, 4.35), (6.78, 4.35)
    ]

    for (title, desc, color), (x, y) in zip(cards, coords):
        add_card(slide, x, y, card_w, card_h, color)
        tb = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.18), Inches(card_w - 0.4), Inches(card_h - 0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_MAIN
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 2)


def build_slide_3(prs):
    """SLIDE 3: CÔNG NGHỆ NHÓM ĐÃ SỬ DỤNG"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(
        slide,
        "03 // CÔNG NGHỆ CÀI ĐẶT",
        "Những Công Nghệ Chính Nhóm Đã Sử Dụng",
        "Lựa chọn các công nghệ hiện đại, ổn định và tối ưu để xây dựng ứng dụng web chạy nhanh và mượt mà"
    )

    # 4 Tech Cards (2x2 grid)
    techs = [
        ("🎨 Giao Diện Web (Frontend)",
         "HTML5, CSS Neo-Futuristic & Chart.js",
         COLOR_NEON_CYAN,
         [
             "• Giao diện phong cách Neon tối giản, hiện đại và rất bắt mắt.",
             "• Thao tác mượt mà, đổi trang ngay lập tức không cần tải lại.",
             "• Biểu đồ tròn và biểu đồ cột trực quan, dễ nhìn cơ cấu chi tiêu."
         ]),
        ("⚡ Xử Lý Hệ Thống (Backend)",
         "Ngôn ngữ Python & Khung FastAPI",
         COLOR_NEON_EMERALD,
         [
             "• Sử dụng Python với FastAPI cho tốc độ xử lý siêu nhanh.",
             "• Hệ thống chạy ổn định, an toàn và dễ dàng mở rộng.",
             "• Đăng nhập an toàn bằng mã khóa JWT bảo vệ 24 giờ."
         ]),
        ("🧠 Trí Tuệ Nhân Tạo (AI Engine)",
         "Google Gemini 1.5 Pro",
         COLOR_NEON_PINK,
         [
             "• Hiểu tiếng Việt cực tốt, nhận diện từ lóng và cách viết tắt.",
             "• Tự động phân loại chi tiêu và đưa ra lời khuyên tài chính 24/7.",
             "• Tốc độ phản hồi cực nhanh, chỉ mất khoảng 0.5 giây."
         ]),
        ("🗄️ Cơ Sở Dữ Liệu (Database)",
         "Hệ quản trị CSDL SQLite",
         COLOR_NEON_AMBER,
         [
             "• Thiết kế 11 bảng dữ liệu khoa học, lưu trữ thông tin gọn gàng.",
             "• Phân chia rõ ràng: ví ghi chép hàng ngày và ví tiền thanh toán.",
             "• Lưu trữ dữ liệu an toàn, không bao giờ lo mất thông tin giao dịch."
         ])
    ]

    card_w = 5.75
    card_h = 2.25
    coords = [
        (0.8, 1.85), (6.78, 1.85),
        (0.8, 4.35), (6.78, 4.35)
    ]

    for (title, sub, color, points), (x, y) in zip(techs, coords):
        add_card(slide, x, y, card_w, card_h, color)
        tb = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(card_w - 0.4), Inches(card_h - 0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(12.5)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_s = tf.add_paragraph()
        p_s.text = f"Công nghệ: {sub}"
        p_s.font.name = FONT_MAIN
        p_s.font.size = Pt(9.5)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_TEXT_MUTED

        for pt in points:
            p_p = tf.add_paragraph()
            p_p.text = pt
            p_p.font.name = FONT_MAIN
            p_p.font.size = Pt(10)
            p_p.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 3)


def build_slide_4(prs):
    """SLIDE 4: TÍNH NĂNG NỔI BẬT 1 - NHẬP NHANH BẰNG CÂU NÓI VỚI AI"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(
        slide,
        "04 // TÍNH NĂNG NỔI BẬT 1",
        "Nhập Chi Tiêu Siêu Nhanh Bằng Câu Nói Với AI",
        "Không cần bấm chọn từng ô phức tạp - Chỉ cần gõ hoặc nói một câu tự nhiên như nhắn tin cho bạn bè"
    )

    # Left Card: Cách thức hoạt động của AI (60% width)
    add_card(slide, 0.8, 1.85, 6.8, 4.85, COLOR_NEON_CYAN)
    ltb = slide.shapes.add_textbox(Inches(1.05), Inches(2.0), Inches(6.3), Inches(4.5))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = "⚡ CÁCH AI TỰ ĐỘNG HIỂU & LƯU GIAO DỊCH"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_CYAN

    p_in = ltf.add_paragraph()
    p_in.text = "💬 Ví dụ câu gõ: \"Ăn sáng bún bò 35k ví MoMo\" hoặc \"Lương tháng này 15tr vào MB Bank\""
    p_in.font.name = FONT_MAIN
    p_in.font.size = Pt(10.5)
    p_in.font.bold = True
    p_in.font.color.rgb = COLOR_NEON_EMERALD

    steps = [
        ("1. Tự Hiểu Ngôn Ngữ Tự Nhiên:",
         "Người dùng gõ câu nói bình thường với từ viết tắt như 'k', 'tr', 'lít', 'củ'. AI vẫn hiểu chính xác 100%."),
        ("2. Tự Động Bóc Tách Đầy Đủ 4 Thông Tin:",
         "• Số tiền: Tự quy đổi 35k ➜ 35.000đ | 15tr ➜ 15.000.000đ.\n• Phân loại: Tự biết là Tiền Ăn Uống (Chi tiêu) hay Tiền Lương (Thu nhập).\n• Ví thanh toán: Tự gán vào ví MoMo, MB Bank hoặc Tiền mặt."),
        ("3. Tốc Độ Phản Hồi Tức Thì:",
         "Xong ngay chỉ sau 0.4 - 0.5 giây, giảm tới 90% thời gian so với cách nhập tay truyền thống.")
    ]
    for st_title, st_desc in steps:
        p_st = ltf.add_paragraph()
        p_st.text = st_title
        p_st.font.name = FONT_HEADING
        p_st.font.size = Pt(11)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_TEXT_WHITE

        p_sd = ltf.add_paragraph()
        p_sd.text = st_desc
        p_sd.font.name = FONT_MAIN
        p_sd.font.size = Pt(10)
        p_sd.font.color.rgb = COLOR_TEXT_LIGHT

    # Right Card: Bảo vệ thông tin nhạy cảm (40% width)
    add_card(slide, 7.8, 1.85, 4.733, 4.85, COLOR_NEON_PINK)
    rtb = slide.shapes.add_textbox(Inches(8.05), Inches(2.0), Inches(4.25), Inches(4.5))
    rtf = rtb.text_frame
    rtf.word_wrap = True

    p = rtf.paragraphs[0]
    p.text = "🛡️ TỰ ĐỘNG XÓA DỮ LIỆU NHẠY CẢM"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_PINK

    p_sub = rtf.add_paragraph()
    p_sub.text = "Cam kết bảo vệ quyền riêng tư tuyệt đối cho người dùng"
    p_sub.font.name = FONT_MAIN
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = COLOR_TEXT_MUTED

    pii_items = [
        ("Nỗi Lo Của Người Dùng:",
         "Sợ bị lộ số tài khoản ngân hàng, số thẻ ngân hàng hoặc tên thật khi gửi dữ liệu lên AI."),
        ("Cách Xử Lý Của FinTrack AI:",
         "• Hệ thống tự động quét và che đi mọi số tài khoản ngân hàng, số thẻ trước khi gửi sang máy chủ AI.\n• Chỉ gửi nội dung chi tiêu thuần túy để AI phân tích."),
        ("Hiệu Quả Bảo Vệ:",
         "Thông tin cá nhân được giữ an toàn 100% trên máy của người dùng, không bao giờ bị rò rỉ ra ngoài.")
    ]
    for p_title, p_desc in pii_items:
        p_pt = rtf.add_paragraph()
        p_pt.text = p_title
        p_pt.font.name = FONT_HEADING
        p_pt.font.size = Pt(11)
        p_pt.font.bold = True
        p_pt.font.color.rgb = COLOR_TEXT_WHITE

        p_pd = rtf.add_paragraph()
        p_pd.text = p_desc
        p_pd.font.name = FONT_MAIN
        p_pd.font.size = Pt(9.8)
        p_pd.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 4)


def build_slide_5(prs):
    """SLIDE 5: TÍNH NĂNG NỔI BẬT 2 - CỐ VẤN TÀI CHÍNH 50/30/20 & RÈN LUYỆN THÓI QUEN"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(
        slide,
        "05 // TÍNH NĂNG NỔI BẬT 2",
        "Trợ Lý Cố Vấn 50/30/20 & Rèn Luyện Thói Quen Tiết Kiệm",
        "Hướng dẫn chia tiền khoa học và biến việc ghi chép sổ sách thành trò chơi thú vị mỗi ngày"
    )

    # Left Card: Cố Vấn 50/30/20 (50% width)
    add_card(slide, 0.8, 1.85, 5.75, 4.85, COLOR_NEON_AMBER)
    ltb = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.35), Inches(4.5))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = "🩺 TRỢ LÝ CỐ VẤN TÀI CHÍNH 50/30/20"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_AMBER

    items_advisor = [
        ("Quy Tắc Vàng 50/30/20 Rất Dễ Áp Dụng:",
         "• 50% Nhu cầu thiết yếu: Tiền ăn uống, tiền nhà, tiền xăng xe, điện nước.\n• 30% Sở thích cá nhân: Mua sắm quần áo, đi xem phim, cà phê bạn bè.\n• 20% Tiết kiệm & Tích lũy: Để dành cho trường hợp khẩn cấp hoặc tương lai."),
        ("AI Tự Động Phân Tích & Cảnh Báo:",
         "• Tổng kết thu chi hàng tuần, hàng tháng xem người dùng tiêu đúng tỷ lệ chưa.\n• Phát chuông cảnh báo khi tiền ăn chơi, mua sắm vượt quá 30% thu nhập."),
        ("Đưa Ra Lời Khuyên Cụ Thể:",
         "Gợi ý chi tiết các khoản nên cắt giảm để tháng sau không bị thâm hụt ngân sách.")
    ]
    for a_title, a_desc in items_advisor:
        p_t = ltf.add_paragraph()
        p_t.text = a_title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_WHITE

        p_d = ltf.add_paragraph()
        p_d.text = a_desc
        p_d.font.name = FONT_MAIN
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = COLOR_TEXT_LIGHT

    # Right Card: Hệ thống rèn luyện thói quen (50% width)
    add_card(slide, 6.78, 1.85, 5.75, 4.85, COLOR_NEON_EMERALD)
    rtb = slide.shapes.add_textbox(Inches(6.98), Inches(2.0), Inches(5.35), Inches(4.5))
    rtf = rtb.text_frame
    rtf.word_wrap = True

    p = rtf.paragraphs[0]
    p.text = "🎮 RÈN LUYỆN THÓI QUEN NHƯ CHƠI GAME"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_EMERALD

    items_game = [
        ("🔥 Chuỗi Ngày Chăm Chỉ (Streak):",
         "Đếm số ngày liên tục người dùng ghi chép chi tiêu. Càng ghi đều đặn, ngọn lửa chuỗi ngày càng cháy to để tạo động lực duy trì thói quen."),
        ("⭐ Tích Lũy Điểm Thưởng & Lên Cấp:",
         "Mỗi lần ghi chép hoặc tiết kiệm thành công được cộng điểm kinh nghiệm (XP), thăng hạng từ 'Tập Sự' lên 'Bậc Thầy Tiết Kiệm'."),
        ("🏆 Bộ 24 Huy Hiệu Thành Tích Đẹp Mắt:",
         "• Huy hiệu 'Chiến Thần Tiết Kiệm': Đạt mục tiêu tiết kiệm 3 tháng liền.\n• Huy hiệu 'Người Tiêu Dùng Thông Thái': Không bao giờ bội chi.\n• Mở khóa các huy hiệu tạo cảm giác tự hào và thích thú.")
    ]
    for g_title, g_desc in items_game:
        p_t = rtf.add_paragraph()
        p_t.text = g_title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_WHITE

        p_d = rtf.add_paragraph()
        p_d.text = g_desc
        p_d.font.name = FONT_MAIN
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 5)


def build_slide_6(prs):
    """SLIDE 6: TÍNH NĂNG NỔI BẬT 3 - HỆ THỐNG 2 LOẠI VÍ TIỀN RÕ RÀNG"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(
        slide,
        "06 // TÍNH NĂNG NỔI BẬT 3",
        "Mô Hình 2 Loại Ví Tiền: Ví Ảo Ghi Chép & Ví Tiền Thật Mua Gói",
        "Phân chia rành mạch để người dùng vừa ghi chép thoải mái, vừa nạp tiền mua dịch vụ an toàn"
    )

    # 2 Big Side-by-Side Cards
    # Card 1: Ví Kế Toán (Ví Ảo)
    add_card(slide, 0.8, 1.85, 5.75, 4.25, COLOR_NEON_CYAN)
    ltb = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.35), Inches(3.95))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = "🌐 1. VÍ KẾ TOÁN (VÍ ẢO GHI CHÉP THU CHI)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_CYAN

    items_v1 = [
        ("Mục đích:", "Dùng để theo dõi tiền tiêu hàng ngày của người dùng."),
        ("Các ví tự tạo:", "Ví MB Bank, Ví MoMo, Ví Tiền Mặt, Thẻ Ngân Hàng..."),
        ("Đặc điểm:", "Số dư ảo do người dùng tự nhập và sửa theo ý muốn. Có thể tạo thêm ví mới hoặc chuyển tiền qua lại giữa các ví ảo."),
        ("Ý nghĩa:", "Hoàn toàn không liên quan đến tiền thật, người dùng thoải mái thử nghiệm ghi chép mà không lo mất tiền.")
    ]
    for label, val in items_v1:
        p_l = ltf.add_paragraph()
        p_l.text = f"• {label} "
        p_l.font.size = Pt(10.5)
        p_l.font.bold = True
        p_l.font.color.rgb = COLOR_TEXT_WHITE
        
        run_v = p_l.add_run()
        run_v.text = val
        run_v.font.bold = False
        run_v.font.color.rgb = COLOR_TEXT_LIGHT

    # Card 2: Ví Tiền Thật
    add_card(slide, 6.78, 1.85, 5.75, 4.25, COLOR_NEON_AMBER)
    rtb = slide.shapes.add_textbox(Inches(6.98), Inches(2.0), Inches(5.35), Inches(3.95))
    rtf = rtb.text_frame
    rtf.word_wrap = True

    p = rtf.paragraphs[0]
    p.text = "💎 2. VÍ TIỀN THẬT (DÙNG ĐỂ MUA GÓI NÂNG CAO)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_AMBER

    items_v2 = [
        ("Mục đích:", "Dùng để thanh toán khi người dùng muốn nâng cấp gói Pro, VIP."),
        ("Cách nạp tiền:", "Chuyển khoản thật qua mã VietQR ngân hàng."),
        ("Đặc điểm:", "Số dư tiền thật được bảo mật tuyệt đối, có hóa đơn điện tử và lịch sử nạp rõ ràng từng đồng."),
        ("Ý nghĩa:", "Đảm bảo tính minh bạch, người dùng nạp bao nhiêu tiền sẽ hiển thị chính xác bấy nhiêu để mua các gói cước.")
    ]
    for label, val in items_v2:
        p_l = rtf.add_paragraph()
        p_l.text = f"• {label} "
        p_l.font.size = Pt(10.5)
        p_l.font.bold = True
        p_l.font.color.rgb = COLOR_TEXT_WHITE
        
        run_v = p_l.add_run()
        run_v.text = val
        run_v.font.bold = False
        run_v.font.color.rgb = COLOR_TEXT_LIGHT

    # Bottom Summary Bar
    bot_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.22), Inches(11.733), Inches(0.55)
    )
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = COLOR_BG_BADGE
    bot_card.line.color.rgb = COLOR_NEON_EMERALD
    bot_card.line.width = Pt(1.2)
    tf_b = bot_card.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "✦ TÓM LẠI: Tách riêng 2 loại ví giúp dữ liệu ghi chép cá nhân không bị lẫn lộn với tiền nạp mua gói, đảm bảo an toàn tuyệt đối."
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.name = FONT_MAIN
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_TEXT_WHITE

    add_footer(slide, 6)


def build_slide_7(prs):
    """SLIDE 7: TÍNH NĂNG NỔI BẬT 4 - NÂNG CẤP & GIA HẠN GÓI CƯỚC TỰ ĐỘNG"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(
        slide,
        "07 // TÍNH NĂNG NỔI BẬT 4",
        "Nâng Cấp & Gia Hạn Gói Cước Tự Động Trong 3 Giây",
        "4 cấp độ gói cước phù hợp mọi nhu cầu và 2 cách thanh toán linh hoạt, tiện lợi"
    )

    # Top 4 Mini Cards for 4 Plans
    plans = [
        ("GÓI MIỄN PHÍ", "0đ", "(Free)", COLOR_TEXT_MUTED, [
            "• 30 lượt AI / tháng",
            "• 1 ví ghi chép cơ bản",
            "• Báo cáo thu chi chuẩn",
            "• Dành cho người mới"
        ]),
        ("GÓI PRO", "49.000đ", "/ tháng", COLOR_NEON_CYAN, [
            "• 200 lượt AI / tháng",
            "• Mở nhiều ví thoải mái",
            "• Cố vấn tài chính 50/30/20",
            "• Biểu đồ phân tích chi tiết"
        ]),
        ("GÓI PREMIUM", "99.000đ", "/ tháng", COLOR_NEON_PURPLE, [
            "• Dùng AI không giới hạn",
            "• Dự báo chi tiêu tương lai",
            "• Xuất file Excel / PDF",
            "• Hỗ trợ ưu tiên 24/7"
        ]),
        ("PLATINUM VIP", "199.000đ", "/ tháng", COLOR_NEON_AMBER, [
            "• Trọn bộ tính năng cao nhất",
            "• Trợ lý AI chuyên sâu",
            "• Giao diện Neon Gold VIP",
            "• Giảm 20% khi mua gói năm"
        ])
    ]

    card_w = 2.76
    spacing = 0.23
    for i, (name, price, period, col, feats) in enumerate(plans):
        cx = 0.8 + i * (card_w + spacing)
        add_card(slide, cx, 1.85, card_w, 2.3, col)
        
        tb = slide.shapes.add_textbox(Inches(cx + 0.1), Inches(1.95), Inches(card_w - 0.2), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_n = tf.paragraphs[0]
        p_n.text = name
        p_n.alignment = PP_ALIGN.CENTER
        p_n.font.name = FONT_HEADING
        p_n.font.size = Pt(11.5)
        p_n.font.bold = True
        p_n.font.color.rgb = col

        p_p = tf.add_paragraph()
        p_p.text = f"{price} {period}"
        p_p.alignment = PP_ALIGN.CENTER
        p_p.font.name = FONT_HEADING
        p_p.font.size = Pt(12.5)
        p_p.font.bold = True
        p_p.font.color.rgb = COLOR_TEXT_WHITE

        for feat in feats:
            p_f = tf.add_paragraph()
            p_f.text = feat
            p_f.font.name = FONT_MAIN
            p_f.font.size = Pt(9)
            p_f.font.color.rgb = COLOR_TEXT_LIGHT

    # Bottom Section: 2 Cách Thanh Toán Linh Hoạt
    add_card(slide, 0.8, 4.3, 11.733, 2.4, COLOR_NEON_EMERALD, COLOR_BG_CARD_LIGHT)
    btb = slide.shapes.add_textbox(Inches(1.0), Inches(4.45), Inches(11.333), Inches(2.1))
    btf = btb.text_frame
    btf.word_wrap = True

    p_bt = btf.paragraphs[0]
    p_bt.text = "💳 2 CÁCH GIA HẠN GÓI SIÊU TIỆN LỢI"
    p_bt.font.name = FONT_HEADING
    p_bt.font.size = Pt(12.5)
    p_bt.font.bold = True
    p_bt.font.color.rgb = COLOR_NEON_EMERALD

    ways = [
        ("👉 Cách 1: Trừ trực tiếp số dư Ví Tiền Thật:",
         "Nếu trong ví đã có sẵn tiền nạp, chỉ cần bấm 'Gia Hạn', tài khoản được nâng cấp ngay lập tức trong 0.1 giây."),
        ("👉 Cách 2: Quét mã VietQR ngân hàng (Tự động nhận diện sau 3 giây):",
         "Chỉ cần mở app ngân hàng quét mã QR, hệ thống tự nhận tiền và tự nâng cấp tài khoản ngay trên màn hình mà KHÔNG CẦN TẢI LẠI TRANG.")
    ]
    for w_title, w_desc in ways:
        p_wt = btf.add_paragraph()
        p_wt.text = w_title
        p_wt.font.name = FONT_HEADING
        p_wt.font.size = Pt(11)
        p_wt.font.bold = True
        p_wt.font.color.rgb = COLOR_TEXT_WHITE

        p_wd = btf.add_paragraph()
        p_wd.text = f"   {w_desc}"
        p_wd.font.name = FONT_MAIN
        p_wd.font.size = Pt(10)
        p_wd.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 7)


def build_slide_8(prs):
    """SLIDE 8: TRANG QUẢN TRỊ DÀNH CHO ADMIN"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(
        slide,
        "08 // TRANG QUẢN TRỊ ADMIN",
        "Trang Quản Trị Hệ Thống Dành Cho Admin",
        "Theo dõi toàn bộ hoạt động của trang web, quản lý người dùng và doanh thu một cách dễ dàng"
    )

    # 3 Admin Function Cards
    cards = [
        ("📊 1. Theo Dõi Tổng Quan",
         COLOR_NEON_CYAN,
         [
             ("Doanh thu thực tế:", "Thống kê chính xác số tiền thu được từ các đơn nạp tiền đã thanh toán thành công."),
             ("Số lượng tài khoản:", "Biết được có bao nhiêu người dùng đang hoạt động và tỷ lệ người dùng gói VIP."),
             ("Mức dùng AI:", "Theo dõi số câu hỏi AI được gửi lên để kiểm soát tài nguyên hệ thống.")
         ]),
        ("⚡ 2. Quản Lý Người Dùng & Đơn Nạp",
         COLOR_NEON_EMERALD,
         [
             ("Khóa tài khoản 1-Click:", "Khóa ngay tài khoản vi phạm hoặc có dấu hiệu gian lận chỉ bằng một nút bấm."),
             ("Duyệt đơn nạp nhanh:", "Xem lại toàn bộ lịch sử quét mã VietQR và trạng thái cộng tiền của từng tài khoản."),
             ("Lịch sử minh bạch:", "Lưu rõ ràng ngày giờ nạp tiền để dễ dàng đối chiếu khi cần thiết.")
         ]),
        ("⚙️ 3. Tùy Chỉnh Lời Nhắc AI",
         COLOR_NEON_AMBER,
         [
             ("Đổi câu lệnh cho AI:", "Trực tiếp sửa hướng dẫn cho trợ lý AI ngay trên web mà không cần khởi động lại máy chủ."),
             ("Nhật ký bảo mật:", "Ghi lại mọi thao tác quan trọng của Admin như đổi mật khẩu hay duyệt tiền."),
             ("An tâm vận hành:", "Đảm bảo trang web luôn hoạt động trơn tru và an toàn 24/7.")
         ])
    ]

    card_w = 3.65
    spacing = 0.39
    for i, (title, color, items) in enumerate(cards):
        cx = 0.8 + i * (card_w + spacing)
        add_card(slide, cx, 1.85, card_w, 4.85, color)
        
        tb = slide.shapes.add_textbox(Inches(cx + 0.15), Inches(2.0), Inches(card_w - 0.3), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = color

        for item_title, item_desc in items:
            p_it = tf.add_paragraph()
            p_it.text = f"✦ {item_title}"
            p_it.font.name = FONT_HEADING
            p_it.font.size = Pt(10.5)
            p_it.font.bold = True
            p_it.font.color.rgb = COLOR_TEXT_WHITE

            p_id = tf.add_paragraph()
            p_id.text = item_desc
            p_id.font.name = FONT_MAIN
            p_id.font.size = Pt(9.8)
            p_id.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 8)


def build_slide_9(prs):
    """SLIDE 9: KẾT QUẢ KIỂM THỬ ỨNG DỤNG"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(
        slide,
        "09 // KẾT QUẢ KIỂM THỬ",
        "Kết Quả Kiểm Thử Thực Tế & Tốc Độ Ứng Dụng",
        "Đã kiểm tra kỹ lưỡng toàn bộ tính năng và đảm bảo trang web hoạt động hoàn hảo 100%"
    )

    # Left Column: 16 Kịch Bản Kiểm Thử (50% width)
    add_card(slide, 0.8, 1.85, 5.75, 4.85, COLOR_NEON_EMERALD)
    ltb = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.35), Inches(4.5))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = "🧪 16 KỊCH BẢN KIỂM THỬ THỰC TẾ"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_EMERALD

    test_groups = [
        ("1. Đăng Ký & Đăng Nhập:",
         "Đăng ký tài khoản mới, lưu mật khẩu an toàn, đăng nhập nhận mã bảo mật đúng quy trình."),
        ("2. Ghi Chép Bằng Câu Nói AI:",
         "Nhập thử nhiều câu nói tiếng Việt phức tạp, AI đều hiểu đúng và lưu chuẩn xác vào danh mục."),
        ("3. Cảnh Báo Tiêu Quá Tay & 50/30/20:",
         "Thử nhập chi tiêu vượt ngân sách, hệ thống lập tức phát cảnh báo đỏ nhắc nhở kịp thời."),
        ("4. Nạp Tiền & Gia Hạn Gói Tự Động:",
         "Quét mã QR chuyển khoản thử, hệ thống tự động cộng tiền và nâng cấp VIP sau 3 giây.")
    ]
    for g_title, g_desc in test_groups:
        p_gt = ltf.add_paragraph()
        p_gt.text = g_title
        p_gt.font.name = FONT_HEADING
        p_gt.font.size = Pt(10.5)
        p_gt.font.bold = True
        p_gt.font.color.rgb = COLOR_TEXT_WHITE

        p_gd = ltf.add_paragraph()
        p_gd.text = g_desc
        p_gd.font.name = FONT_MAIN
        p_gd.font.size = Pt(9.5)
        p_gd.font.color.rgb = COLOR_TEXT_LIGHT

    # Right Column: Kết Quả & Tốc Độ (50% width)
    add_card(slide, 6.78, 1.85, 5.75, 4.85, COLOR_NEON_CYAN)
    rtb = slide.shapes.add_textbox(Inches(6.98), Inches(2.0), Inches(5.35), Inches(4.5))
    rtf = rtb.text_frame
    rtf.word_wrap = True

    p = rtf.paragraphs[0]
    p.text = "🚀 KẾT QUẢ ĐẠT ĐƯỢC & HIỆU NĂNG"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_CYAN

    perf_items = [
        ("🏆 KẾT QUẢ KIỂM THỬ TỰ ĐỘNG:", 
         "100% CÁC BÀI TEST ĐẠT KẾT QUẢ CHÍNH XÁC (PASSED)", 
         COLOR_NEON_EMERALD),
        ("⚡ TỐC ĐỘ AI NHẬN DIỆN TIẾNG VIỆT:", 
         "Chỉ mất khoảng 0.4 - 0.5 giây để bóc tách xong một câu nói.", 
         COLOR_NEON_CYAN),
        ("⏱️ TỐC ĐỘ TẢI TRANG WEB:", 
         "Dưới 0.1 giây nhờ công nghệ FastAPI hiện đại.", 
         COLOR_NEON_AMBER),
        ("🎨 TRẢI NGHIỆM HÌNH ẢNH MƯỢT MÀ:", 
         "Hiệu ứng viền phát sáng Neon chạy mượt mà đạt chuẩn 120 khung hình/giây, không bị giật lag.", 
         COLOR_NEON_PURPLE)
    ]
    for title, val, col in perf_items:
        p_t = rtf.add_paragraph()
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_WHITE

        p_v = rtf.add_paragraph()
        p_v.text = f"✦ {val}"
        p_v.font.name = FONT_MAIN
        p_v.font.size = Pt(9.5)
        p_v.font.bold = True
        p_v.font.color.rgb = col

    add_footer(slide, 9)


def build_slide_10(prs):
    """SLIDE 10: ĐÁNH GIÁ & HƯỚNG PHÁT TRIỂN"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(
        slide,
        "10 // ĐÁNH GIÁ & PHÁT TRIỂN",
        "Đánh Giá Kết Quả Đạt Được & Hướng Phát Triển Tiếp Theo",
        "Nhìn lại những gì nhóm đã làm được và các tính năng dự định mở rộng trong tương lai"
    )

    # Left Card: Ưu điểm đã làm được (50% width)
    add_card(slide, 0.8, 1.85, 5.75, 4.85, COLOR_NEON_EMERALD)
    ltb = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.35), Inches(4.5))
    ltf = ltb.text_frame
    ltf.word_wrap = True

    p = ltf.paragraphs[0]
    p.text = "⭐ KẾT QUẢ NỔI BẬT ĐÃ ĐẠT ĐƯỢC"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_EMERALD

    strengths = [
        ("✦ Hoàn Thiện Một Trang Web Đầy Đủ:",
         "Ứng dụng chạy thực tế rất mượt mà, đầy đủ các chức năng từ ghi chép, biểu đồ đến nạp tiền tự động."),
        ("✦ AI Hiểu Tiếng Việt Cực Tốt:",
         "Nhập liệu bằng câu nói tự nhiên rất nhanh và chính xác, tự động xóa số tài khoản ngân hàng để bảo mật."),
        ("✦ Mô Hình 2 Ví Tiền Rõ Ràng:",
         "Tách riêng ví ghi chép ảo và ví nạp tiền thật, không bị lẫn lộn dữ liệu."),
        ("✦ Quét Mã QR Tự Động Hóa 100%:",
         "Tự nhận tiền chuyển khoản sau 3 giây và tự nâng cấp gói VIP mà không cần làm thủ công.")
    ]
    for s_title, s_desc in strengths:
        p_st = ltf.add_paragraph()
        p_st.text = s_title
        p_st.font.name = FONT_HEADING
        p_st.font.size = Pt(11)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_TEXT_WHITE

        p_sd = ltf.add_paragraph()
        p_sd.text = s_desc
        p_sd.font.name = FONT_MAIN
        p_sd.font.size = Pt(9.8)
        p_sd.font.color.rgb = COLOR_TEXT_LIGHT

    # Right Card: Hướng phát triển tương lai (50% width)
    add_card(slide, 6.78, 1.85, 5.75, 4.85, COLOR_NEON_AMBER)
    rtb = slide.shapes.add_textbox(Inches(6.98), Inches(2.0), Inches(5.35), Inches(4.5))
    rtf = rtb.text_frame
    rtf.word_wrap = True

    p = rtf.paragraphs[0]
    p.text = "🚀 HƯỚNG PHÁT TRIỂN TIẾP THEO"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_AMBER

    roadmap = [
        ("🔮 1. Làm Ứng Dụng Trên Điện Thoại (App Mobile):",
         "Phát triển thêm phiên bản app cho điện thoại Android và iPhone để người dùng ghi chép mọi lúc mọi nơi tiện lợi hơn."),
        ("🔮 2. Chụp Ảnh Hóa Đơn Tự Động Lưu Sổ (OCR):",
         "Chỉ cần giơ máy ảnh chụp hóa đơn ăn uống, đi siêu thị, AI sẽ tự đọc từng món hàng và lưu vào sổ chi tiêu."),
        ("🔮 3. Kết Nối Trực Tiếp Với Ngân Hàng:",
         "Tự động nhận thông báo biến động số dư từ tài khoản ngân hàng để không bao giờ quên ghi chép.")
    ]
    for r_title, r_desc in roadmap:
        p_rt = rtf.add_paragraph()
        p_rt.text = r_title
        p_rt.font.name = FONT_HEADING
        p_rt.font.size = Pt(11)
        p_rt.font.bold = True
        p_rt.font.color.rgb = COLOR_TEXT_WHITE

        p_rd = rtf.add_paragraph()
        p_rd.text = r_desc
        p_rd.font.name = FONT_MAIN
        p_rd.font.size = Pt(9.8)
        p_rd.font.color.rgb = COLOR_TEXT_LIGHT

    add_footer(slide, 10)


def build_slide_11(prs):
    """SLIDE 11: LỜI CẢM ƠN & PHẦN HỎI ĐÁP (Q&A)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    # Top Tag
    top_badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.4), Inches(0.55), Inches(6.533), Inches(0.4)
    )
    top_badge.fill.solid()
    top_badge.fill.fore_color.rgb = COLOR_BG_BADGE
    top_badge.line.color.rgb = COLOR_NEON_EMERALD
    top_badge.line.width = Pt(1)
    tf_badge = top_badge.text_frame
    p_b = tf_badge.paragraphs[0]
    p_b.text = "✦ BÁO CÁO KẾT THÚC ĐỀ TÀI HỌC PHẦN AI ✦"
    p_b.alignment = PP_ALIGN.CENTER
    p_b.font.name = FONT_MAIN
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_NEON_EMERALD

    # Big Q&A Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.733), Inches(1.25))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True

    p1 = tf_t.paragraphs[0]
    p1.text = "XIN TRÂN TRỌNG CẢM ƠN CÔ GIÁO & CÁC BẠN!"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.name = FONT_HEADING
    p1.font.size = Pt(27)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_WHITE

    p2 = tf_t.add_paragraph()
    p2.text = "PHIÊN HỎI ĐÁP & GÓP Ý ĐỀ TÀI (Q&A SESSION)"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(17)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_NEON_CYAN

    # 3 Member Contact Cards
    members = [
        ("✦ Đặng Quyết Thắng", "Trưởng Nhóm / Phát Triển AI & Web", "dangquyetthang@ictu.edu.vn", COLOR_NEON_CYAN),
        ("✦ Nguyễn Văn Tiến", "Thành Viên / Thiết Kế Giao Diện Web", "nguyenvantien@ictu.edu.vn", COLOR_NEON_EMERALD),
        ("✦ Quách Minh Hiếu", "Thành Viên / Xử Lý CSDL & Backend", "quachminhhieu@ictu.edu.vn", COLOR_NEON_PURPLE)
    ]
    card_w = 3.65
    spacing = 0.39
    for i, (name, role, email, col) in enumerate(members):
        cx = 0.8 + i * (card_w + spacing)
        add_card(slide, cx, 2.45, card_w, 2.2, col)
        
        tb = slide.shapes.add_textbox(Inches(cx + 0.15), Inches(2.6), Inches(card_w - 0.3), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = name
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_WHITE

        p_r = tf.add_paragraph()
        p_r.text = role
        p_r.alignment = PP_ALIGN.CENTER
        p_r.font.name = FONT_MAIN
        p_r.font.size = Pt(10)
        p_r.font.bold = True
        p_r.font.color.rgb = col

        p_e = tf.add_paragraph()
        p_e.text = f"Email: {email}"
        p_e.alignment = PP_ALIGN.CENTER
        p_e.font.name = FONT_MAIN
        p_e.font.size = Pt(9.5)
        p_e.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom Live Demo Banner Card
    bot_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.85), Inches(11.733), Inches(1.8)
    )
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = COLOR_BG_CARD_LIGHT
    bot_card.line.color.rgb = COLOR_NEON_AMBER
    bot_card.line.width = Pt(1.5)
    
    btb = slide.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(11.333), Inches(1.5))
    btf = btb.text_frame
    btf.word_wrap = True

    p = btf.paragraphs[0]
    p.text = "✦ NHÓM ĐÃ SẴN SÀNG CHẠY THỬ ỨNG DỤNG TRỰC TIẾP (LIVE DEMO) ✦"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_HEADING
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_NEON_AMBER

    p_d1 = btf.add_paragraph()
    p_d1.text = "• Trình diễn Gõ một câu nói tự nhiên để AI tự lưu giao dịch siêu tốc trong 0.5 giây."
    p_d1.alignment = PP_ALIGN.CENTER
    p_d1.font.name = FONT_MAIN
    p_d1.font.size = Pt(10.5)
    p_d1.font.color.rgb = COLOR_TEXT_LIGHT

    p_d2 = btf.add_paragraph()
    p_d2.text = "• Trình diễn Quét mã VietQR chuyển khoản thật và tự động nâng cấp VIP sau 3 giây."
    p_d2.alignment = PP_ALIGN.CENTER
    p_d2.font.name = FONT_MAIN
    p_d2.font.size = Pt(10.5)
    p_d2.font.color.rgb = COLOR_TEXT_LIGHT

    p_d3 = btf.add_paragraph()
    p_d3.text = "Kính mời Cô giáo và các bạn đặt câu hỏi thảo luận cùng nhóm!"
    p_d3.alignment = PP_ALIGN.CENTER
    p_d3.font.name = FONT_MAIN
    p_d3.font.size = Pt(11)
    p_d3.font.bold = True
    p_d3.font.color.rgb = COLOR_NEON_CYAN

    # Footer for Slide 11
    add_footer(slide, 11, total_slides=11, accent_color=COLOR_NEON_EMERALD)


# ==============================================================================
# HÀM CHÍNH KHỞI TẠO VÀ XUẤT FILE PPTX
# ==============================================================================

def main():
    print("[INFO] Đang khởi tạo bản thuyết trình FinTrack AI phiên bản tiếng Việt tự nhiên...")
    prs = create_presentation()

    print("[INFO] Đang tạo Slide 1: Trang Bìa...")
    build_slide_1(prs)

    print("[INFO] Đang tạo Slide 2: Vì Sao Nhóm Chọn Đề Tài Này?...")
    build_slide_2(prs)

    print("[INFO] Đang tạo Slide 3: Công Nghệ Nhóm Đã Sử Dụng...")
    build_slide_3(prs)

    print("[INFO] Đang tạo Slide 4: Tính Năng 1 - Nhập Nhanh Bằng Câu Nói Với AI...")
    build_slide_4(prs)

    print("[INFO] Đang tạo Slide 5: Tính Năng 2 - Cố Vấn Tài Chính 50/30/20 & Rèn Luyện Thói Quen...")
    build_slide_5(prs)

    print("[INFO] Đang tạo Slide 6: Tính Năng 3 - Hệ Thống 2 Loại Ví Tiền Rõ Ràng...")
    build_slide_6(prs)

    print("[INFO] Đang tạo Slide 7: Tính Năng 4 - Nâng Cấp & Gia Hạn Gói Cước Tự Động...")
    build_slide_7(prs)

    print("[INFO] Đang tạo Slide 8: Trang Quản Trị Dành Cho Admin...")
    build_slide_8(prs)

    print("[INFO] Đang tạo Slide 9: Kết Quả Kiểm Thử Ứng Dụng...")
    build_slide_9(prs)

    print("[INFO] Đang tạo Slide 10: Đánh Giá & Hướng Phát Triển...")
    build_slide_10(prs)

    print("[INFO] Đang tạo Slide 11: Lời Cảm Ơn & Phần Hỏi Đáp (Q&A)...")
    build_slide_11(prs)

    output_filename = "BaoCao_FinTrackAI_Demo.pptx"
    output_path = os.path.abspath(output_filename)

    try:
        prs.save(output_path)
        print(f"[SUCCESS] XUẤT THÀNH CÔNG! File báo cáo đã được lưu tại: {output_path}")
    except PermissionError:
        fallback_filename = "BaoCao_FinTrackAI_ThuyetTrinh_Moi.pptx"
        fallback_path = os.path.abspath(fallback_filename)
        prs.save(fallback_path)
        print(f"[WARNING] File '{output_filename}' đang được mở trong PowerPoint!")
        print(f"[SUCCESS] Đã lưu bản thuyết trình mới tại: {fallback_path}")
        print(f"[TIP] Vui lòng đóng PowerPoint hoặc mở file '{fallback_filename}' để xem.")

    print(f"[SUCCESS] Tổng số slide đã tạo: {len(prs.slides)}")

if __name__ == "__main__":
    main()
